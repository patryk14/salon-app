"""Automatic Booksy pull (F5) — download the same report Booksy hands out as
xlsx and feed it through the EXISTING, tested xlsx parser. No second parser.

Booksy has no public API; this uses its internal business API the way the
browser does, authenticating with the three headers captured from a session
(x-access-token / x-api-key / x-fingerprint). The access token rotates (~daily),
so credentials live in a DB row the owner refreshes via an admin endpoint —
never in code or git. Uses stdlib urllib (a plain GET returning bytes) so the
runtime image gains no dependency.
"""

import io
import json
import urllib.parse
import urllib.request
from dataclasses import dataclass

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.booksy import parse_visits_report
from app.models import BooksyCredential
from app.routers.imports import ImportSummary, import_visit_rows

REPORT_URL = (
    "https://pl.booksy.com/core/v2/business_api/me/stats/businesses/{business_id}/report/download"
)


class BooksyAuthError(RuntimeError):
    """No credentials configured, or Booksy rejected them (expired token)."""


@dataclass(frozen=True)
class BooksyCredentials:
    business_id: str
    access_token: str
    api_key: str
    fingerprint: str


def load_credentials(db: Session) -> BooksyCredentials:
    row = db.scalar(select(BooksyCredential).order_by(BooksyCredential.id).limit(1))
    if row is None:
        raise BooksyAuthError("brak zapisanych poświadczeń Booksy — ustaw token")
    return BooksyCredentials(
        business_id=row.business_id,
        access_token=row.access_token,
        api_key=row.api_key,
        fingerprint=row.fingerprint,
    )


def download_report(
    creds: BooksyCredentials,
    report_key: str,
    date_from: str,
    date_till: str,
    time_span: str = "month",
    timeout: int = 90,
) -> bytes:
    """GET the report/download endpoint → xlsx bytes. Mirrors the browser's
    request headers; a 401/403 means the session token has expired. The
    cash-registers report returns full history and is slow to generate, hence
    the generous default timeout."""
    params = urllib.parse.urlencode(
        {
            "date_from": date_from,
            "date_till": date_till,
            "time_span": time_span,
            "report_key": report_key,
        }
    )
    url = f"{REPORT_URL.format(business_id=creds.business_id)}?{params}"
    req = urllib.request.Request(
        url,
        headers={
            "x-access-token": creds.access_token,
            "x-api-key": creds.api_key,
            "x-fingerprint": creds.fingerprint,
            "x-app-version": "3.0",
            "accept": "application/json, text/plain, */*",
            "user-agent": "charmskin-sync/1.0",
        },
    )
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:  # noqa: S310 (fixed https host)
            return resp.read()
    except urllib.error.HTTPError as e:
        if e.code in (401, 403):
            raise BooksyAuthError("Booksy odrzucił poświadczenia — token wygasł, odśwież") from e
        raise BooksyAuthError(f"Booksy zwrócił HTTP {e.code}") from e


CUSTOMERS_URL = "https://pl.booksy.com/core/v2/business_api/me/businesses/{business_id}/customers"


def _get_json(creds: BooksyCredentials, url: str, timeout: int = 60) -> dict:
    """Authenticated GET returning JSON (same header shape as download_report)."""
    req = urllib.request.Request(
        url,
        headers={
            "x-access-token": creds.access_token,
            "x-api-key": creds.api_key,
            "x-fingerprint": creds.fingerprint,
            "x-app-version": "3.0",
            "accept": "application/json, text/plain, */*",
            "user-agent": "charmskin-sync/1.0",
        },
    )
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:  # noqa: S310 (fixed https host)
            return json.loads(resp.read())
    except urllib.error.HTTPError as e:
        if e.code in (401, 403):
            raise BooksyAuthError("Booksy odrzucił poświadczenia — token wygasł, odśwież") from e
        raise BooksyAuthError(f"Booksy zwrócił HTTP {e.code}") from e


def pull_customers(db: Session, per_page: int = 100, max_pages: int = 60) -> dict:
    """Backfill client contacts + consents from Booksy's customers API (F7 v2).

    Booksy's `id` is the stable identity key; `merged_data` carries name, phone,
    email and consents. Match an existing Client by that id, else by name (only a
    row not already linked, so we never steal another person's link), else create
    a new one. Contacts/consents are the base for self-signup and reminders."""
    from app.models import Client, ClientTombstone

    creds = load_credentials(db)
    existing = db.scalars(select(Client)).all()
    # RODO tombstones: Booksy ids of clients erased on request — never re-import.
    tombstoned = set(db.scalars(select(ClientTombstone.booksy_customer_id)).all())
    by_booksy = {c.booksy_customer_id: c for c in existing if c.booksy_customer_id}
    by_name: dict[tuple[str, str], list] = {}
    for c in existing:
        by_name.setdefault((c.first_name.strip().lower(), c.last_name.strip().lower()), []).append(
            c
        )

    total = created = updated = with_email = with_phone = 0
    url_base = CUSTOMERS_URL.format(business_id=creds.business_id)
    for page in range(1, max_pages + 1):
        data = _get_json(creds, f"{url_base}?per_page={per_page}&page={page}")
        customers = data.get("customers", [])
        if not customers:
            break
        for cu in customers:
            md = cu.get("merged_data") or cu.get("customer_profile") or {}
            bid = md.get("id") or cu.get("_id")
            first = (md.get("first_name") or "").strip()
            last = (md.get("last_name") or "").strip()
            if not (first or last):  # nameless ghost record → skip
                continue
            if bid and bid in tombstoned:  # RODO-erased → do not recreate or re-link
                continue
            total += 1
            phone = (md.get("cell_phone") or "").strip() or None
            email = (md.get("email") or "").strip() or None

            row = by_booksy.get(bid) if bid else None
            if row is None:
                candidates = [
                    c
                    for c in by_name.get((first.lower(), last.lower()), [])
                    if c.booksy_customer_id is None
                ]
                row = candidates[0] if candidates else None
            if row is None:
                row = Client(first_name=first or "?", last_name=last or "?")
                db.add(row)
                created += 1
            else:
                updated += 1
            if bid:
                row.booksy_customer_id = bid
                by_booksy[bid] = row
            if phone:
                row.phone = phone
                with_phone += 1
            if email:
                row.email = email
                with_email += 1
            row.marketing_consent = bool(md.get("marketing_agreement"))
            row.privacy_consent = bool(md.get("privacy_policy_agreement"))
        db.flush()
        if len(customers) < per_page:
            break
    return {
        "customers": total,
        "created": created,
        "updated": updated,
        "with_email": with_email,
        "with_phone": with_phone,
    }


def pull_visits(db: Session, report_key: str, date_from: str, date_till: str) -> ImportSummary:
    """Download the visits report for a date range and upsert it — the same
    path as the manual xlsx upload, just fetched automatically."""
    creds = load_credentials(db)
    data = download_report(creds, report_key, date_from, date_till)
    rows = parse_visits_report(io.BytesIO(data))
    return import_visit_rows(db, rows)


def pull_registers(db: Session, date_from: str, date_till: str) -> dict:
    """Derive per-day salon-till figures (gotówka z Booksy + kasa fiskalna) from
    Booksy's cash-registers TRANSACTIONS report into salon_days. That report
    respects the URL date range (unlike the 'summary'), matches the owner's sheet
    to the złoty, and flags package redemptions. Authoritative for those two
    figures — replaces the range's salon_day rows (the ledger, i.e. 'gotówka nie
    wbita', lives elsewhere and is untouched)."""
    from datetime import date as _date

    from openpyxl import load_workbook
    from sqlalchemy import delete

    from app.models import SalonDay
    from app.registers import parse_cash_transactions

    creds = load_credentials(db)
    data = download_report(creds, "cash_registers_transactions", date_from, date_till)
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    grid = [list(r) for r in wb.active.iter_rows(values_only=True)]
    parsed = parse_cash_transactions(grid)

    start, end = _date.fromisoformat(date_from), _date.fromisoformat(date_till)
    days = {d: v for d, v in parsed.by_day.items() if start <= d <= end}
    db.execute(delete(SalonDay).where(SalonDay.day >= start, SalonDay.day <= end))
    for d, v in days.items():
        db.add(SalonDay(day=d, booksy_cash=v["booksy_cash"], fiscal_register=v["fiscal_register"]))
    db.flush()
    red = _sync_package_redemptions(db, parsed.package_txs, start, end)
    cash_total = sum((v["booksy_cash"] for v in days.values()), start=0)
    fiscal_total = sum((v["fiscal_register"] for v in days.values()), start=0)
    return {
        "days": len(days),
        "sessions": parsed.transactions,
        "package_redemptions": parsed.package_redemptions,
        "redemptions_matched": red["matched"],
        "redemptions_unmatched": red["unmatched"],
        "cash_total": str(cash_total),
        "fiscal_total": str(fiscal_total),
    }


def _sync_package_redemptions(db: Session, package_txs: list, start, end) -> dict:
    """For each 'Pakiet' till transaction: find the client's package active that
    day (one active per client → unambiguous) and the performer from that day's
    visit (transaction is always the shared cashier), then upsert a redemption
    crediting value_per_treatment. Idempotent on the Booksy document number;
    package/employee left NULL (flagged as unmatched) when they can't be
    resolved — never silently dropped."""
    from datetime import timedelta
    from decimal import Decimal

    from app.booksy import split_name
    from app.models import Client, EmployeeAlias, Package, PackageRedemption, Visit

    alias_map = dict(db.execute(select(EmployeeAlias.alias, EmployeeAlias.employee_id)).all())
    matched = unmatched = 0
    for d, client_name, doc in package_txs:
        if not (start <= d <= end) or not doc:
            continue
        first, last = split_name(client_name)
        client = db.scalar(
            select(Client).where(Client.first_name == first, Client.last_name == last)
        )
        pkg = emp_id = None
        if client is not None:
            pkg = db.scalar(
                select(Package)
                .where(
                    Package.client_id == client.id,
                    Package.valid_from <= d,
                    Package.valid_until >= d,
                )
                .order_by(Package.valid_until)
            )
            staff = db.scalar(
                select(Visit.staff_name)
                .where(
                    Visit.client_id == client.id,
                    Visit.starts_at >= d,
                    Visit.starts_at < d + timedelta(days=1),
                )
                .limit(1)
            )
            if staff:
                emp_id = alias_map.get(staff)
        value = (
            (pkg.total_value / pkg.total_treatments).quantize(Decimal("0.01"))
            if pkg
            else Decimal("0")
        )
        row = db.scalar(select(PackageRedemption).where(PackageRedemption.booksy_ref == doc))
        if row is None:
            row = PackageRedemption(booksy_ref=doc)
            db.add(row)
        row.package_id = pkg.id if pkg else None
        row.employee_id = emp_id
        row.client_name = client_name
        row.redemption_date = d
        row.value = value
        if pkg is not None and emp_id is not None:
            matched += 1
        else:
            unmatched += 1
    db.flush()
    return {"matched": matched, "unmatched": unmatched}


def pull_packages(db: Session) -> dict:
    """Sync every client package from Booksy's packages_summary (Booksy is the
    source — packages are sold there). Upserts by Booksy package number; matches
    the client by name (link only, never creates). The report ignores dates and
    returns the full current state, which is what we want."""
    from datetime import date

    from openpyxl import load_workbook
    from sqlalchemy import func

    from app.booksy import split_name
    from app.models import Client, Package, PackageRedemption
    from app.packages import parse_packages_summary

    creds = load_credentials(db)
    data = download_report(creds, "packages_summary", "2026-01-01", "2026-12-31")
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    grid = [list(r) for r in wb.active.iter_rows(values_only=True)]
    parsed = parse_packages_summary(grid)

    # Owner ruling (2026-09-14): packages that expired before the current year are
    # "dead" and no longer tracked. Don't import them, and self-prune any stale
    # rows left from an earlier sync (only where no redemption history depends on
    # them, so we never orphan a credited commission).
    cutoff = date(date.today().year, 1, 1)

    created = updated = active = skipped_expired = 0
    for p in parsed:
        if p.valid_until is not None and p.valid_until < cutoff:
            skipped_expired += 1
            continue
        first, last = split_name(p.client_name)
        client = db.scalar(
            select(Client).where(Client.first_name == first, Client.last_name == last)
        )
        row = db.scalar(select(Package).where(Package.booksy_number == p.booksy_number))
        if row is None:
            row = Package(booksy_number=p.booksy_number)
            db.add(row)
            created += 1
        else:
            updated += 1
        row.client_id = client.id if client else None
        row.client_name = p.client_name
        row.name = p.name
        row.total_value = p.total_value
        row.total_treatments = p.total_treatments
        row.remaining = p.remaining
        row.valid_from = p.valid_from
        row.valid_until = p.valid_until
        if p.remaining > 0:
            active += 1

    removed_expired = 0
    for row in db.scalars(select(Package).where(Package.valid_until < cutoff)).all():
        linked = db.scalar(
            select(func.count())
            .select_from(PackageRedemption)
            .where(PackageRedemption.package_id == row.id)
        )
        if not linked:
            db.delete(row)
            removed_expired += 1

    db.flush()
    return {
        "packages": len(parsed),
        "created": created,
        "updated": updated,
        "active": active,
        "skipped_expired": skipped_expired,
        "removed_expired": removed_expired,
    }
