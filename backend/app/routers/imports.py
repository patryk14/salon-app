"""Booksy xlsx import.

Idempotent by design: `booksy_ref` (Booksy reservation id) is unique in the
visits table, so re-uploading the same or an overlapping report UPDATES
existing rows instead of duplicating them. Clients are matched by display name
(the only identity Booksy reports expose) and created on first sight.
"""

import logging
from decimal import Decimal
from typing import Annotated

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from pydantic import BaseModel, Field
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.auth import require_role
from app.booksy import BooksyParseError, VisitRow, parse_visits_report, split_name
from app.costs import parse_costs
from app.deps import get_db
from app.derivation import month_bounds
from app.models import (
    BooksyCredential,
    Client,
    Employee,
    EmployeeAlias,
    LedgerEntry,
    SalonDay,
    Visit,
)

logger = logging.getLogger(__name__)

# Imports mutate the whole dataset (creates clients, rewrites visits) — owner only.
router = APIRouter(prefix="/imports", tags=["imports"], dependencies=[require_role("admin")])


class ImportSummary(BaseModel):
    visits_in_file: int
    clients_created: int
    visits_created: int
    visits_updated: int


class CostImportSummary(BaseModel):
    year_month: str
    ledger_created: int
    salon_days_created: int
    per_employee: dict[str, str]  # employee name -> month cash sum (zł)
    unmatched_names: list[str]  # sheet rows no alias resolved — cash dropped
    checksum_ok: bool  # Σ our cash per day == the sheet's "Gotówka nie wbita" row?
    checksum_note: str


class BooksyCredentialsIn(BaseModel):
    business_id: str = Field(default="221497", max_length=20)
    access_token: str = Field(min_length=1, max_length=200)
    api_key: str = Field(min_length=1, max_length=200)
    fingerprint: str = Field(min_length=1, max_length=200)


class PullRequest(BaseModel):
    # report_key identifies WHICH Booksy report to download (the visits list).
    report_key: str = Field(min_length=1, max_length=100)
    date_from: str = Field(pattern=r"^\d{4}-\d{2}-\d{2}$")
    date_till: str = Field(pattern=r"^\d{4}-\d{2}-\d{2}$")


def _client_for(db: Session, cache: dict[str, Client], name: str) -> tuple[Client, bool]:
    if name in cache:
        return cache[name], False
    first, last = split_name(name)
    existing = db.scalars(
        select(Client).where(Client.first_name == first, Client.last_name == last)
    ).first()
    if existing is not None:
        cache[name] = existing
        return existing, False
    client = Client(first_name=first, last_name=last)
    db.add(client)
    db.flush()
    cache[name] = client
    return client, True


def _upsert_visit(db: Session, client: Client, row: VisitRow) -> bool:
    """Returns True when a new visit was created, False on update."""
    visit = db.scalars(select(Visit).where(Visit.booksy_ref == row.booksy_ref)).first()
    if visit is None:
        db.add(
            Visit(
                client_id=client.id,
                booksy_ref=row.booksy_ref,
                starts_at=row.starts_at,
                service_name=row.service_name,
                status=row.status,
                price_pln=row.price_pln,
                staff_name=row.staff_name,
            )
        )
        return True
    # Booksy is the source of truth for its own fields (a rebooked or cancelled
    # visit must win over our copy); local-only fields (notes) stay untouched.
    visit.client_id = client.id
    visit.starts_at = row.starts_at
    visit.service_name = row.service_name
    visit.status = row.status
    visit.price_pln = row.price_pln
    visit.staff_name = row.staff_name
    return False


def import_visit_rows(db: Session, rows: list[VisitRow]) -> ImportSummary:
    """Upsert parsed visit rows — the shared core of both the xlsx upload and the
    automatic pull (F5). Idempotent on booksy_ref."""
    cache: dict[str, Client] = {}
    clients_created = visits_created = visits_updated = 0
    for row in rows:
        client, created = _client_for(db, cache, row.client_name)
        clients_created += created
        if _upsert_visit(db, client, row):
            visits_created += 1
        else:
            visits_updated += 1

    summary = ImportSummary(
        visits_in_file=len(rows),
        clients_created=clients_created,
        visits_created=visits_created,
        visits_updated=visits_updated,
    )
    logger.info("booksy import: %s", summary.model_dump())
    return summary


@router.post("/booksy/visits", status_code=status.HTTP_200_OK)
def import_booksy_visits(
    db: Annotated[Session, Depends(get_db)],
    file: Annotated[UploadFile, File(description="Booksy Biz 'Lista wizyt' xlsx export")],
) -> ImportSummary:
    try:
        rows = parse_visits_report(file.file)
    except BooksyParseError as e:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, detail=str(e)) from e
    return import_visit_rows(db, rows)


@router.post("/costs", status_code=status.HTTP_200_OK)
def import_costs(
    db: Annotated[Session, Depends(get_db)],
    year_month: Annotated[str, Form(pattern=r"^\d{4}-(0[1-9]|1[0-2])$")],
    file: Annotated[UploadFile, File(description="zabiegi_koszty monthly till xlsx")],
) -> CostImportSummary:
    """Backfill a month's off-Booksy cash from the zabiegi_koszty sheet. The
    sheet has no year/month, so the caller names the month. REPLACES that month's
    ledger + salon_day rows (idempotent re-import; use for historical months, not
    a month being entered daily)."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(file.file, read_only=True, data_only=True)
        grid = [list(r) for r in wb.active.iter_rows(values_only=True)]
    except Exception as e:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, detail="unreadable xlsx") from e

    # Resolve sheet names (e.g. "Karolina") to employees via aliases + display names.
    alias_map = dict(db.execute(select(EmployeeAlias.alias, EmployeeAlias.employee_id)).all())
    name_map = dict(db.execute(select(Employee.display_name, Employee.id)).all())
    names = dict(db.execute(select(Employee.id, Employee.display_name)).all())

    def resolve(s: str) -> int | None:
        return alias_map.get(s) or name_map.get(s)

    year, month = (int(p) for p in year_month.split("-"))
    parsed = parse_costs(grid, year, month, resolve)

    # Replace the month: this import is authoritative for its cash.
    start, end = month_bounds(year_month)
    db.execute(
        delete(LedgerEntry).where(LedgerEntry.entry_date >= start, LedgerEntry.entry_date < end)
    )
    db.execute(delete(SalonDay).where(SalonDay.day >= start, SalonDay.day < end))
    for emp_id, d, amt in parsed.ledger:
        db.add(
            LedgerEntry(
                employee_id=emp_id, entry_date=d, service_name="Gotówka (arkusz)", amount_pln=amt
            )
        )
    salon_days_created = 0
    for d, vals in parsed.salon_days.items():
        db.add(
            SalonDay(
                day=d,
                booksy_cash=vals.get("booksy_cash", Decimal("0")),
                fiscal_register=vals.get("fiscal_register", Decimal("0")),
            )
        )
        salon_days_created += 1
    db.flush()

    # Checksum: our Σ ledger per day vs the sheet's own "Gotówka nie wbita" row.
    by_day: dict = {}
    for _emp, d, amt in parsed.ledger:
        by_day[d] = by_day.get(d, Decimal("0")) + amt
    mismatched = [
        d for d, exp in parsed.unregistered_by_day.items() if by_day.get(d, Decimal("0")) != exp
    ]
    if mismatched:
        days = [d.isoformat() for d in sorted(mismatched)[:5]]
        note = f"rozjazd w {len(mismatched)} dniach: {days}"
    else:
        note = "suma gotówki zgadza się z wierszem kontrolnym arkusza"
    summary = CostImportSummary(
        year_month=year_month,
        ledger_created=len(parsed.ledger),
        salon_days_created=salon_days_created,
        per_employee={names.get(eid, f"#{eid}"): str(s) for eid, s in parsed.per_employee.items()},
        unmatched_names=sorted(set(parsed.unmatched_names)),
        checksum_ok=not mismatched,
        checksum_note=note,
    )
    logger.info("costs import %s: %s", year_month, summary.model_dump())
    return summary


@router.put("/booksy/credentials", status_code=status.HTTP_204_NO_CONTENT)
def set_booksy_credentials(
    payload: BooksyCredentialsIn, db: Annotated[Session, Depends(get_db)]
) -> None:
    """Store/refresh the Booksy session credentials for the automatic pull. The
    access token rotates — the owner pastes a fresh one here, no redeploy."""
    row = db.scalars(select(BooksyCredential).order_by(BooksyCredential.id).limit(1)).first()
    if row is None:
        row = BooksyCredential(id=1)
        db.add(row)
    row.business_id = payload.business_id
    row.access_token = payload.access_token
    row.api_key = payload.api_key
    row.fingerprint = payload.fingerprint


@router.post("/booksy/pull", status_code=status.HTTP_200_OK)
def pull_booksy_visits(
    payload: PullRequest, db: Annotated[Session, Depends(get_db)]
) -> ImportSummary:
    """Download the visits report for a date range straight from Booksy and
    upsert it — automation of the manual xlsx upload."""
    from app.booksy_api import BooksyAuthError, pull_visits  # lazy: avoids import cycle

    try:
        return pull_visits(db, payload.report_key, payload.date_from, payload.date_till)
    except BooksyAuthError as e:
        raise HTTPException(status.HTTP_502_BAD_GATEWAY, detail=str(e)) from e
    except BooksyParseError as e:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, detail=str(e)) from e
